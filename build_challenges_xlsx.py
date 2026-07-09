"""Editable workbook of all challenges. One row per challenge, every field a
column. ID column maps back to the seed order for re-import after edits."""
import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment

challenges = json.load(open('challenges.json', encoding='utf-8'))

wb = Workbook()
ws = wb.active
ws.title = 'Challenges'

GREEN = '1B4332'
GREENMID = '2D6A4F'
GREY = 'F3F3F1'

cols = [
    ('ID', 6), ('Category', 16), ('Title', 30), ('Difficulty (1-3)', 11),
    ('Time (min)', 10), ('Cost', 10), ('Description', 42),
    ('Instructions (one per line)', 42), ('Why it matters', 46),
    ('Tips (one per line)', 34), ('Barriers (one per line)', 28),
    ('CO2 (kg)', 9), ('Water (L)', 9), ('Waste (kg)', 9), ('Region', 9),
]

hdr_fill = PatternFill('solid', fgColor=GREEN)
hdr_font = Font(name='Arial', bold=True, color='FFFFFF', size=10)
base_font = Font(name='Arial', size=10)
id_font = Font(name='Arial', size=10, color='999999')
thin = Side(style='thin', color='D9D9D6')
border = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap_top = Alignment(wrap_text=True, vertical='top')
center = Alignment(horizontal='center', vertical='top')

# header
for i, (name, width) in enumerate(cols, 1):
    cell = ws.cell(row=1, column=i, value=name)
    cell.fill = hdr_fill; cell.font = hdr_font
    cell.alignment = Alignment(vertical='center', wrap_text=True)
    cell.border = border
    ws.column_dimensions[cell.column_letter].width = width
ws.cell(row=1, column=1).comment = Comment(
    'Do not edit the ID column. It maps each row back to the source for re-import.', 'Casuarina')

def joinlist(v):
    return '\n'.join(v) if isinstance(v, list) else (v or '')

for r, c in enumerate(challenges, start=2):
    ie = c.get('impactEstimate', {}) or {}
    row = [
        r - 1, c.get('category',''), c.get('title',''), c.get('difficulty',''),
        c.get('timeEstimate',''), c.get('costLevel',''), c.get('description',''),
        joinlist(c.get('instructions')), c.get('educationalText',''),
        joinlist(c.get('tips')), joinlist(c.get('barriers')),
        ie.get('co2', 0), ie.get('water', 0), ie.get('waste', 0), c.get('region',''),
    ]
    for i, val in enumerate(row, 1):
        cell = ws.cell(row=r, column=i, value=val)
        cell.font = id_font if i == 1 else base_font
        cell.border = border
        cell.alignment = center if i in (1,4,5,6,12,13,14,15) else wrap_top
    if r % 2 == 0:
        for i in range(1, len(cols)+1):
            if i != 1:
                ws.cell(row=r, column=i).fill = PatternFill('solid', fgColor=GREY)

n = len(challenges) + 1
# dropdowns
dv_diff = DataValidation(type='list', formula1='"1,2,3"', allow_blank=False)
dv_cost = DataValidation(type='list', formula1='"free,low,medium,high"', allow_blank=False)
ws.add_data_validation(dv_diff); ws.add_data_validation(dv_cost)
dv_diff.add(f'D2:D{n}')
dv_cost.add(f'F2:F{n}')

ws.freeze_panes = 'D2'
ws.auto_filter.ref = f'A1:O{n}'
ws.sheet_view.showGridLines = True

wb.save('challenges.xlsx')
print('built challenges.xlsx with', len(challenges), 'rows')
