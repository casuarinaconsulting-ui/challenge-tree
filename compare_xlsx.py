import json
from openpyxl import load_workbook

XLSX = r'C:\Users\ramon\casuarina-ai\challenge-tre3-challenges.xlsx'
orig = json.load(open('challenges.json', encoding='utf-8'))

wb = load_workbook(XLSX)
ws = wb.active
rows = list(ws.iter_rows(min_row=2, values_only=True))

def splitlines(v):
    if v is None: return []
    return [s.strip() for s in str(v).split('\n') if s.strip()]

def norm(v):
    return '' if v is None else str(v).strip()

edited = []
issues = []
COSTS = {'free','low','medium','high'}
for r in rows:
    if r[0] is None and all(c is None for c in r):
        continue
    (ID, cat, title, diff, time, cost, desc, instr, why, tips, barr, co2, water, waste, region) = r
    e = {
        'ID': ID, 'category': norm(cat), 'title': norm(title),
        'difficulty': diff, 'timeEstimate': time, 'costLevel': norm(cost),
        'description': norm(desc), 'instructions': splitlines(instr),
        'educationalText': norm(why), 'tips': splitlines(tips),
        'barriers': splitlines(barr),
        'impactEstimate': {'co2': co2, 'water': water, 'waste': waste},
        'region': norm(region) or 'global',
    }
    edited.append(e)
    # validations
    if not e['title']:
        issues.append(f"ID {ID}: empty title")
    if e['costLevel'] not in COSTS:
        issues.append(f"ID {ID} '{e['title']}': invalid cost '{e['costLevel']}'")
    if e['difficulty'] not in (1,2,3):
        issues.append(f"ID {ID} '{e['title']}': invalid difficulty '{e['difficulty']}'")
    for k in ('co2','water','waste'):
        v = e['impactEstimate'][k]
        if not isinstance(v,(int,float)):
            issues.append(f"ID {ID} '{e['title']}': impact {k} not numeric: {v!r}")
    if not isinstance(e['timeEstimate'],(int,float)):
        issues.append(f"ID {ID} '{e['title']}': time not numeric: {e['timeEstimate']!r}")

print(f"rows in sheet: {len(edited)}  | original: {len(orig)}")

by_id = {e['ID']: e for e in edited}
orig_ids = set(range(1, len(orig)+1))
edit_ids = set(by_id.keys())
new_rows = [e for e in edited if e['ID'] not in orig_ids]
missing = sorted(orig_ids - edit_ids)
if new_rows: print(f"NEW rows (ID not in 1..{len(orig)}): {[ (e['ID'],e['title']) for e in new_rows]}")
if missing: print(f"MISSING IDs (deleted rows): {missing}")

FIELDS = ['category','title','difficulty','timeEstimate','costLevel','description',
          'instructions','educationalText','tips','barriers','region']
changes = []
for i, o in enumerate(orig, start=1):
    e = by_id.get(i)
    if not e: continue
    for f in FIELDS:
        ov, ev = o.get(f), e.get(f)
        if isinstance(ov, list): ov = [x.strip() for x in ov]
        if ov != ev:
            changes.append((i, o['title'], f, ov, ev))
    # impact compare
    oie = o.get('impactEstimate',{}) or {}
    eie = e['impactEstimate']
    for k in ('co2','water','waste'):
        if (oie.get(k) or 0) != (eie.get(k) or 0):
            changes.append((i, o['title'], f'impact.{k}', oie.get(k), eie.get(k)))

print(f"\n=== {len(changes)} field change(s) ===")
for (i,t,f,ov,ev) in changes:
    print(f"\nID {i} [{t}] . {f}")
    print(f"  OLD: {ov!r}")
    print(f"  NEW: {ev!r}")

print(f"\n=== {len(issues)} validation issue(s) ===")
for s in issues: print(" -", s)

json.dump(edited, open('challenges_edited.json','w',encoding='utf-8'), ensure_ascii=False, indent=2)
